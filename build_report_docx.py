from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "REPORT.md"
OUTPUT = ROOT / "Лабораторная_работа_1_Гарифуллин_Искандар_P4108.docx"

NAVY = "17324D"
BLUE = "DCEAF4"
PALE = "F5F8FA"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, font_size=11, bold=False, italic=False):
    """Write markdown-ish inline text while retaining links and emphasis."""
    pattern = re.compile(r"(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.font.size = Pt(font_size)
            run.bold = bold
            run.italic = italic
        token = match.group(0)
        if token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            is_bold = token.startswith("**")
            inner = token[2:-2] if is_bold else token[1:-1]
            run = paragraph.add_run(inner)
            run.font.size = Pt(font_size)
            run.bold = bold or is_bold
            run.italic = italic or (not is_bold)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.first_line_indent = Inches(0.35)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, space_before, space_after in [
        ("Title", 22, 0, 14),
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 10, 4),
    ]:
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLACK
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.keep_with_next = True

    if "Report Caption" not in styles:
        caption = styles.add_style("Report Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Arial"
        caption.font.size = Pt(10)
        caption.font.bold = True
        caption.font.color.rgb = BLACK
        caption.paragraph_format.space_before = Pt(6)
        caption.paragraph_format.space_after = Pt(4)
    if "Table Text" not in styles:
        table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
        table_text.font.name = "Arial"
        table_text.font.size = Pt(8.6)
        table_text.font.color.rgb = BLACK
        table_text.paragraph_format.space_after = Pt(1)
        table_text.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Лабораторная работа № 1  •  ")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(90, 90, 90)
    add_page_number(footer)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.add_run("УНИВЕРСИТЕТ ИНФОРМАЦИОННЫХ ТЕХНОЛОГИЙ\n").bold = True
    p.add_run("Учебный отчёт по дисциплине\n")
    p.add_run("Методология исследований и разработки ИТ-систем").italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Лабораторная работа № 1")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(52)
    run = subtitle.add_run("Понятия «метод» и «методология» в науке и в решении ИТ-задач")
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.bold = True

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    widths = [2.0, 4.2]
    values = [
        ("Студент", "Гарифуллин Искандар Ильданович"),
        ("Группа", "P4108"),
        ("Тема мини-проекта", "Веб-прототип платформы для постановки и размещения рекламных задач IT CREW DIGITAL"),
        ("Дата", "07.09.2026"),
    ]
    for row, (label, value) in zip(meta.rows, values):
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell, 110, 140, 110, 140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = "Table Text"
            p.paragraph_format.space_after = Pt(0)
            if idx == 0:
                set_cell_shading(cell, NAVY)
                r = p.add_run(label)
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.bold = True
            else:
                set_cell_shading(cell, PALE)
                p.add_run(value)
    set_table_borders(meta)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Репозиторий: ")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    add_hyperlink(p, "github.com/iskandario/lab1-it-crew-digital", "https://github.com/iskandario/lab1-it-crew-digital")
    p.add_run("\n").font.size = Pt(4)
    r = p.add_run("Опубликованный MVP: ")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    add_hyperlink(p, "it-crew-digital.okrupashiani.chatgpt.site", "https://it-crew-digital.okrupashiani.chatgpt.site/")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run("Казань\n2026")
    r.font.name = "Arial"
    r.font.size = Pt(11)

    doc.add_page_break()


def clean_text(text):
    text = text.replace("`", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def add_body_paragraph(doc, text, style=None, indent=True):
    p = doc.add_paragraph(style=style)
    if indent and style is None:
        p.paragraph_format.first_line_indent = Inches(0.35)
    add_inline(p, text)
    return p


def add_list_item(doc, text, ordered=False, font_size=11):
    # Manual numbering keeps each independent numbered list at 1 rather than
    # inheriting Word's global List Number counter from the previous section.
    p = doc.add_paragraph(style=None if ordered else "List Bullet")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(2 if font_size < 11 else 3)
    add_inline(p, text, font_size=font_size)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    n = len(headers)
    if n == 4:
        widths = [1.35, 1.75, 1.85, 2.1]
        font_size = 8.0
    elif n == 3:
        widths = [2.0, 2.45, 2.6]
        font_size = 8.7
    else:
        widths = [1.45] + [5.75 / (n - 1)] * (n - 1)
        font_size = 8.8
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for idx, cell in enumerate(hdr.cells):
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 95, 105, 95, 105)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.style = "Table Text"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(clean_text(headers[idx]))
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.bold = True
    for row_idx, values in enumerate(rows):
        new_row = table.add_row()
        set_row_cant_split(new_row)
        cells = new_row.cells
        for idx, value in enumerate(values):
            set_cell_width(cells[idx], widths[idx])
            set_cell_shading(cells[idx], "FFFFFF" if row_idx % 2 == 0 else PALE)
            set_cell_margins(cells[idx], 90, 105, 90, 105)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.style = "Table Text"
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Inches(0)
            add_inline(p, clean_text(value), font_size=font_size)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_diagram(doc):
    add_table(doc, ["Метод", "ИТ-задача IT CREW DIGITAL", "Методология"], [[
        "конкретное действие или способ проверки",
        "Рекламный бриф и запуск кампании",
        "Итерационная разработка, связывающая методы в цикл улучшений",
    ]])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Метод решает отдельный шаг  →  задача  ←  методология организует весь проект")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.italic = True


def parse_report(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    # Skip the markdown cover because the DOCX has a designed cover page.
    while i < len(lines) and not lines[i].startswith("## Цель работы"):
        i += 1
    table_buffer = None
    mermaid_mode = False
    in_bibliography = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            language = line[3:].strip().lower()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            if language == "mermaid":
                add_diagram(doc)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.space_after = Pt(6)
                for idx, code_line in enumerate(code_lines):
                    r = p.add_run(code_line + ("\n" if idx < len(code_lines) - 1 else ""))
                    r.font.name = "Courier New"
                    r.font.size = Pt(9.5)
            continue
        if line.startswith("|"):
            if table_buffer is None:
                table_buffer = [line]
            else:
                table_buffer.append(line)
            i += 1
            continue
        if table_buffer is not None:
            if len(table_buffer) >= 2:
                raw = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in table_buffer]
                headers = raw[0]
                rows = [r for r in raw[2:] if not all(set(c) <= set("-:") for c in r)]
                add_table(doc, headers, rows)
            table_buffer = None
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(clean_text(line[2:]))
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(clean_text(line[3:]))
            in_bibliography = "Список использованных источников" in line
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(clean_text(line[4:]))
        elif re.match(r"^\d+\.\s+", line):
            add_list_item(doc, line, ordered=True, font_size=10.5 if in_bibliography else 11)
        elif line.startswith("- "):
            add_list_item(doc, line[2:], ordered=False)
        elif line.startswith("**Вывод по таблице.**"):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.35)
            add_inline(p, line)
        else:
            add_body_paragraph(doc, line)
        i += 1
    if table_buffer is not None:
        raw = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in table_buffer]
        add_table(doc, raw[0], [r for r in raw[2:] if not all(set(c) <= set("-:") for c in r)])


def main():
    doc = Document()
    style_document(doc)
    add_cover(doc)
    parse_report(doc)
    doc.core_properties.title = "Лабораторная работа № 1 Метод и методология"
    doc.core_properties.author = "Гарифуллин Искандар Ильданович"
    doc.core_properties.subject = "Методология исследований и веб-прототип IT CREW DIGITAL"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
